import csv, logging
from openpyxl import load_workbook
from .models import Card
from . import utils 
from .utils import format_card, format_expire, format_phone, format_balance, get_live_exchange_rate
from decimal import Decimal, ROUND_HALF_UP



logger = logging.getLogger(__name__)


def export_cards(filters=None, output_file='cards_export.csv'):
    """
    Karta ma'lumotlarini CSV formatda eksport qiladi.

    Args:
        filters (dict, optional): Karta qidiruvi uchun filtr shartlari.
        output_file (str): Chiqariladigan CSV fayl nomi.

    Returns:
        int: Eksport qilingan karta satrlari soni.
    """
    # Bazadagi barcha karta ob'ektlarini QuerySet ko'rinishida oladi
    queryset = Card.objects.all()
    
    # Agar filtrlar uzatilgan bo'lsa, qidiruv shartlarini shakllantiradi
    if filters:
        query_params = {}
        for key, value in filters.items():
            # Status bo'yicha qidiruv aniq (exact) bo'lishi kerak
            if key == 'status':
                query_params[key] = value
            # Boshqa maydonlar (masalan: ism, raqam) uchun qisman moslikni qidiradi
            else:
                query_params[f"{key}__icontains"] = value
        
        # Dinamik ravishda shakllantirilgan filtrni bazaga jo'natadi
        queryset = queryset.filter(**query_params)

    # Faylni yozish rejimida ochadi (UTF-8 formatida)
    with open(output_file, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f) # CSV yozuvchi obyektini yaratadi
        # Jadvalning birinchi qatori - sarlavhalarni yozadi
        writer.writerow(['Karta raqami', 'Muddati', 'Telefon', 'Status', 'Balans', 'Yaratilgan vaqt'])
        
        # Har bir karta bo'yicha sikl aylantiradi va ma'lumotlarni qatorga yozadi
        for card in queryset:
            writer.writerow([
                card.card_number,             # Karta raqami
                card.expire,                  # Amal qilish muddati
                card.phone or '',             # Telefon (bo'sh bo'lsa bo'sh joy)
                card.get_status_display(),    # Statusning tushunarli matni (masalan: 'Faol')
                card.balance,                 # Karta balansi
                card.created_at.strftime("%Y-%m-%d %H:%M") # Vaqtni chiroyli formatda yozish
            ])
            
    return queryset.count() # Jami eksport qilinganlar sonini qaytaradi






logger = logging.getLogger(__name__)

def import_cards(file):
    """
    Excel fayldagi kartalarni bazaga import qiladi.

    Args:
        file: Yuklangan Excel fayl obyekti.

    Returns:
        tuple: (success_count, errors) - muvaffaqiyatli saqlangan kartalar soni va xato xabarlari.
    """
    # Agar fayl yuklanmagan bo'lsa, jarayonni to'xtatadi
    if not file:
        return 0, ["Fayl tanlanmagan"]
    
    try:
        # Excel faylni yuklaydi (faqat formulalarning natijasini oladi)
        wb = load_workbook(file, data_only=True)
        sheet = wb.active # Birinchi (faol) varaqni tanlaydi
        success_count = 0 # Muvaffaqiyatli saqlanganlar hisoblagichi
        errors = []      # Xatolarni yig'ib boruvchi ro'yxat

        # Exceldagi o'zbekcha/inglizcha statuslarni bazadagi formatga o'tkazish xaritasi
        STATUS_MAP = {
            'active': 'active', 'faol': 'active',
            'inactive': 'inactive', 'faol emas': 'inactive',
            'expired': 'expired', 'muddati otgan': 'expired',
        }

        # 2-qatordan boshlab ma'lumotlarni o'qiydi (1-qator sarlavha)
        rows = sheet.iter_rows(min_row=2, values_only=True)

        for row_idx, row in enumerate(rows, start=2):
            raw_card = row[0] # Birinchi ustundagi karta raqami
            # Bo'sh qatorlarni yoki sarlavha takrorlansa tashlab ketadi
            if not raw_card or "card" in str(raw_card).lower():
                continue

            # 1. Ma'lumotlarni utilitalar yordamida tozalash va formatlash
            num = format_card(raw_card)           # Raqamni faqat sonlarga keltirish
            expire = format_expire(row[1])        # Muddati formatini tekshirish (00/00)
            phone = format_phone(row[2])          # Telefon raqamini standartga keltirish
            
            # Statusni aniqlash (bo'sh bo'lsa 'active' deb oladi)
            raw_status = str(row[3]).lower().strip() if row[3] else 'active'
            current_status = STATUS_MAP.get(raw_status, 'active')
            
            # 2. VALYUTA KONVERTATSIYASI
            raw_balance = format_balance(row[4])  # Balansni son formatiga keltirish
            # Valyuta turini aniqlash (Default: UZS)
            currency = str(row[5]).upper().strip() if len(row) > 5 and row[5] else 'UZS'

            # Real vaqtda valyuta kursini oladi
            rate = get_live_exchange_rate(currency)
            # Balansni so'mga ko'paytiradi va 0.01 gacha yaxlitlaydi
            balance_in_uzs = (Decimal(str(raw_balance)) * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

            # Agar karta raqami noto'g'ri bo'lsa xatoni qayd etadi
            if not num:
                errors.append(f"{row_idx}-qatorda karta raqami xato: {raw_card}")
                continue

            try:
                # 3. BAZAGA SAQLASH (Update or Create)
                # Karta raqami bo'yicha qidiradi, bo'lsa yangilaydi, bo'lmasa yaratadi
                Card.objects.update_or_create(
                    card_number=num,
                    defaults={
                        'expire': expire,
                        'phone': phone,
                        'balance': balance_in_uzs, # Doim UZS formatida saqlanadi
                        'status': current_status 
                    }
                )
                success_count += 1 # Hisoblagichni oshiradi
            except Exception as e:
                # Bazaga yozishda xatolik bo'lsa ro'yxatga qo'shadi
                errors.append(f"{row_idx}-qatorda bazaga saqlashda xato: {str(e)}")

        return success_count, errors # Yakuniy natijalarni qaytaradi

    except Exception as e:
        # Excel o'qishda kutilmagan xato bo'lsa
        return 0, [f"Excel o'qishda jiddiy xato: {str(e)}"]




def send_card_messages(status='active', lang='UZ'):
    """
    Filtrlangan kartalarga xabar yuborish xizmatini simulyatsiya qiladi.

    Args:
        status (str): Qaysi holatdagi kartalar uchun xabar yuborilishi kerak.
        lang (str): Xabar tilini belgilaydi (UZ, RU, EN).

    Returns:
        int: Yuborilgan xabarlar soni.
    """
    # Faqat so'ralgan statusdagi kartalarni filtrlab oladi
    cards = Card.objects.filter(status=status)
    
    for card in cards:
        # Har bir karta uchun alohida xabar matni generatsiya qiladi
        msg = utils.prepare_message(card, lang=lang)
        
        # Test uchun terminalga chiqaradi (Production-da bu yerda SMS API bo'ladi)
        print(f"SENDING TO {card.phone}: {msg}")
        
        # Audit loglariga muvaffaqiyatli yuborilganini qayd etadi
        logger.info(f"Sent to {card.card_number}")
        
    return cards.count() # Jami yuborilgan xabarlar sonini qaytaradi